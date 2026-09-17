import os
import sys
import time
import glob
import json
import base64
import requests
import urllib3
import openpyxl
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading

urllib3.disable_warnings()

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = SCRIPT_DIR
DESKTOP_DIR = os.path.join(os.path.expanduser('~'), 'Desktop')

# Auto-detect Excel file
def find_excel_file():
    # 1. Look in current directory first
    for f in glob.glob(os.path.join(SCRIPT_DIR, "*.xlsx")):
        if not os.path.basename(f).startswith("~$"):
            return f
    # 2. Look for EXPORT ON WASEEL W1.xlsx on Desktop
    default_excel = os.path.join(DESKTOP_DIR, "EXPORT ON WASEEL W1.xlsx")
    if os.path.exists(default_excel):
        return default_excel
    # 3. Look for any xlsx on Desktop
    for f in glob.glob(os.path.join(DESKTOP_DIR, "*.xlsx")):
        if not os.path.basename(f).startswith("~$"):
            return f
    return default_excel

EXCEL_PATH = find_excel_file()

def b64decode_padding(s):
    s = s.replace('-', '+').replace('_', '/')
    while len(s) % 4 != 0:
        s += '='
    return base64.b64decode(s)

def get_latest_edge_token():
    edge_profile = os.path.join(os.path.expanduser('~'), r'AppData\Local\Microsoft\Edge\User Data\Profile 1')
    ls_path = os.path.join(edge_profile, 'Local Storage', 'leveldb')
    import re
    jwt_pattern = re.compile(rb'eyJ[a-zA-Z0-9_\-]{20,}\.eyJ[a-zA-Z0-9_\-]{20,}\.[a-zA-Z0-9_\-]+')

    best_token = None
    best_exp = 0

    if os.path.exists(ls_path):
        for fname in os.listdir(ls_path):
            if fname.endswith(('.log', '.ldb')):
                fpath = os.path.join(ls_path, fname)
                try:
                    with open(fpath, 'rb') as f:
                        content = f.read()
                        for m in jwt_pattern.findall(content):
                            t_str = m.decode('ascii', errors='ignore')
                            parts = t_str.split('.')
                            try:
                                payload = json.loads(b64decode_padding(parts[1]))
                                if 'sso.nphies.sa' in payload.get('iss', ''):
                                    exp = payload.get('exp', 0)
                                    if exp > best_exp:
                                        best_exp = exp
                                        best_token = t_str
                            except:
                                pass
                except:
                    pass

    return best_token

class TokenKeeper:
    def __init__(self):
        self.lock = threading.Lock()
        self.token = get_latest_edge_token()

    def get_token(self):
        return self.token

    def refresh(self):
        with self.lock:
            t = get_latest_edge_token()
            if t:
                self.token = t
            return self.token

keeper = TokenKeeper()

def get_headers():
    return {
        "Authorization": f"Bearer {keeper.get_token()}",
        "facilitylicense": "10000000064871",
        "facilitytype": "Provider",
        "username": "10000000064871",
        "origin": "https://viewer.nphies.sa",
        "referer": "https://viewer.nphies.sa/",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36 Edg/151.0.0.0"
    }

def fetch_bundle(b_id, retries=3):
    target = os.path.join(OUTPUT_DIR, f"Claim Request_{b_id}.json")
    if os.path.exists(target) and os.path.getsize(target) > 100:
        return b_id, "Downloaded", os.path.basename(target), round(os.path.getsize(target)/1024, 1), "-"

    for attempt in range(retries):
        try:
            h = get_headers()
            search_url = "https://sgw.nphies.sa/viewerapi/claim"
            params = {
                "size": 10,
                "page": 0,
                "date_from": "2026-07-01T00:00:00.000Z",
                "date_to": "2026-09-15T23:59:59.000Z",
                "bundle_id": b_id
            }
            r = requests.get(search_url, headers=h, params=params, verify=False, timeout=20)
            if r.status_code == 401:
                keeper.refresh()
                time.sleep(1)
                continue
            if r.status_code != 200:
                time.sleep(1)
                continue

            data = r.json()
            items = data.get('Items', [])
            if not items:
                return b_id, "Not Found", "", 0, ""

            res_id = items[0].get('ResourceID')
            tx_id = items[0].get('TransactionID', '')

            detail_url = f"https://sgw.nphies.sa/viewerapi/claim/GetById?id={res_id}&facility_Type=Provider&facility_license=10000000064871"
            r2 = requests.get(detail_url, headers=h, verify=False, timeout=25)
            if r2.status_code == 401:
                keeper.refresh()
                time.sleep(1)
                continue
            if r2.status_code != 200:
                time.sleep(1)
                continue

            detail = r2.json()
            content = detail.get('RequestResourceContent')
            if not content:
                return b_id, "Empty Content", "", 0, tx_id

            with open(target, 'w', encoding='utf-8') as f:
                if isinstance(content, str):
                    f.write(content)
                else:
                    json.dump(content, f, indent=2)

            return b_id, "Downloaded", os.path.basename(target), round(os.path.getsize(target)/1024, 1), tx_id

        except Exception as e:
            time.sleep(1)

    return b_id, "Error", "", 0, ""

def sync_excel():
    if not os.path.exists(EXCEL_PATH):
        return 0, 0, {}

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Find BundleID column
    bundle_col = 1
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(row=1, column=c).value or '').lower().replace(' ', '').replace('_', '')
        if 'bundle' in h:
            bundle_col = c
            break

    ws.cell(row=1, column=2).value = 'Status'
    ws.cell(row=1, column=3).value = 'Transaction_ID'
    ws.cell(row=1, column=4).value = 'File_Name'
    ws.cell(row=1, column=5).value = 'File_Size_KB'
    ws.cell(row=1, column=6).value = 'Downloaded_At'

    bundle_row_map = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=bundle_col).value
        if v:
            bundle_row_map[str(v).strip()] = r

    downloaded = glob.glob(os.path.join(OUTPUT_DIR, '*.json'))
    now_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    matched = 0
    for f in downloaded:
        b_id = os.path.basename(f).replace('Claim Request_', '').replace('.json', '')
        if b_id in bundle_row_map:
            row_idx = bundle_row_map[b_id]
            if ws.cell(row=row_idx, column=2).value != 'Downloaded':
                ws.cell(row=row_idx, column=2).value = 'Downloaded'
                ws.cell(row=row_idx, column=4).value = os.path.basename(f)
                ws.cell(row=row_idx, column=5).value = round(os.path.getsize(f) / 1024, 1)
                ws.cell(row=row_idx, column=6).value = now_str
            matched += 1

    try:
        wb.save(EXCEL_PATH)
    except PermissionError:
        pass

    return matched, len(bundle_row_map), bundle_row_map

def main():
    print('=' * 75)
    print('      NPHIES BULK CLAIM REQUEST DOWNLOADER (run.bat)')
    print('=' * 75)
    print(f'Detected Excel: {EXCEL_PATH}')
    print(f'Save Folder:    {OUTPUT_DIR}')
    print('-' * 75)

    if not os.path.exists(EXCEL_PATH):
        print(f'ERROR: No Excel file found!')
        return

    while True:
        matched, total, bundle_map = sync_excel()
        pct = (matched / total * 100) if total > 0 else 0
        print(f'[{datetime.now().strftime("%H:%M:%S")}] Status: {matched}/{total} ({pct:.1f}%) Downloaded & Synced in Excel')

        if matched >= total and total > 0:
            print('=' * 75)
            print(f'SUCCESS: ALL {total} BUNDLES ARE EXTRACTED AND MATCHED 100% IN EXCEL!')
            print(f'JSON Folder: {OUTPUT_DIR}')
            print(f'Excel File:  {EXCEL_PATH}')
            print('=' * 75)
            break

        downloaded_bundles = set(os.path.basename(f).replace('Claim Request_', '').replace('.json', '') for f in glob.glob(os.path.join(OUTPUT_DIR, '*.json')))
        missing = [b for b in bundle_map.keys() if b not in downloaded_bundles]

        if missing:
            print(f'Fetching batch of missing bundles (Remaining: {len(missing)})...')
            batch = missing[:50]
            with ThreadPoolExecutor(max_workers=4) as ex:
                futs = {ex.submit(fetch_bundle, b): b for b in batch}
                for fut in as_completed(futs):
                    res = fut.result()

        time.sleep(3)

if __name__ == '__main__':
    main()
